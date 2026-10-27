# -*- coding: utf-8 -*-
"""Excel 导出：按质检模板版式输出。

Sheet1「检测数据」——模板版式：
- 第 1-5 行：表头元数据区 12 项（报表名称/工件名称/工件编号/操作者/公司名称/
  车间/工段/楼层/班别/工序/机台号/客户/时间），支持通过 meta 参数传入；
- 第 6 行：空行；
- 第 7 行：表头（序号/元素/内容/名义值/上公差/下公差/测量值）；
- 第 8 行起：每点拆两行（内容列=X坐标/Y坐标），元素命名 P1 为中心点基准、
  P2 缺失、P3 起按行主序映射行列；上/下公差从参数或模板读取（默认 ±0.05）。

名义值来源（固定值优先，禁止检测反推）：
- 传入 nominal_points（模板 D 列固定名义值 26 点 = P1 + P3..P27）时，
  Sheet1 名义值列直接采用固定值；
- 未传入时回退为 rows 中 theo_x/theo_y（设计网格中心归零坐标）。
测量值 = rows 中 mm_x/mm_y（调用方需保证与名义值同坐标系——以网格几何
中心理论位置为原点的中心归零坐标）。

Sheet2「检测明细」：保留像素坐标/ΔX/ΔY/偏差/判定/状态等逐点信息。
Sheet3「参数记录」：检测参数快照。

判定口径由上层负责（X/Y 分量各自落在 [名义+下公差, 名义+上公差]）。
"""
import csv
import math
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.point_order import (MODE_BOTTOM_LEFT_ROW_MAJOR, display_no,
                              display_order)

DEFAULT_TOL_UPPER = 0.05
DEFAULT_TOL_LOWER = -0.05

# ---- B8：8 列判定 CSV 导出（与现有 export_excel 质检模板相互独立） ----
# 判定阈值 0.15mm 写死（项目书 §6），不做配置项；误差 = 名义 − 实测。
DETECT_THRESHOLD_MM = 0.15
DEFAULT_CSV_ORDER = MODE_BOTTOM_LEFT_ROW_MAJOR

CSV_HEADERS_8COL = {
    "zh": ["P", "X名义", "Y名义", "X实测", "Y实测", "X误差", "Y误差", "判定"],
    "en": ["P", "Nominal X", "Nominal Y", "Measured X", "Measured Y",
           "Error X", "Error Y", "Verdict"],
}

_META_ROWS = [
    ("报表名称", "车间/工段", "时间"),
    ("工件名称", "楼层", ""),
    ("工件编号", "班别", ""),
    ("操作者", "工序", ""),
    ("公司名称", "机台号", "客户"),
]

# ---- 导出表头/文案中英双语表（zh 为默认；en 时整表切换语言） ----
_META_LABELS_ZH = ("报表名称", "工件名称", "工件编号", "操作者", "公司名称",
                   "车间/工段", "楼层", "班别", "工序", "机台号", "客户", "时间")
_META_LABELS_EN = ("Report Name", "Workpiece Name", "Workpiece No.", "Operator", "Company",
                   "Workshop/Section", "Floor", "Shift", "Process", "Machine No.", "Customer", "Time")
_SHEET_NAMES = {
    "zh": ("检测数据", "检测明细", "参数记录"),
    "en": ("Inspection Data", "Inspection Details", "Parameters"),
}
_HEADERS = {
    "zh": ["序号", "元素", "内容", "名义值", "上公差", "下公差", "测量值"],
    "en": ["No.", "Element", "Content", "Nominal", "Upper Tol.", "Lower Tol.", "Measured"],
}
_HEADERS2 = {
    "zh": ["序号", "行", "列", "像素X", "像素Y", "实测X(mm)", "实测Y(mm)",
           "理论X(mm)", "理论Y(mm)", "ΔX(mm)", "ΔY(mm)", "偏差(mm)", "判定", "置信度", "状态"],
    "en": ["No.", "Row", "Col", "Pixel X", "Pixel Y", "Measured X(mm)", "Measured Y(mm)",
           "Theoretical X(mm)", "Theoretical Y(mm)", "ΔX(mm)", "ΔY(mm)", "Deviation(mm)", "Verdict", "Confidence", "Status"],
}
_XY_CONTENT = {"zh": ("X坐标", "Y坐标"), "en": ("X Coordinate", "Y Coordinate")}

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_CENTER = Alignment(horizontal="center", vertical="center")


def _is_fail(r: dict) -> bool:
    """失败行判定：status/verdict 含 DETECT_FAIL、失败、未检出 任一即视为失败行。

    兼容改造前的旧数据（status="未检出"），保证统计口径统一排除失败行。
    """
    s = str(r.get("status", ""))
    v = str(r.get("verdict", ""))
    return ("DETECT_FAIL" in s or "失败" in s or "未检出" in s
            or "DETECT_FAIL" in v or "失败" in v or "未检出" in v)


def _center_index(results: List[dict], rows: Optional[int], cols: Optional[int]) -> int:
    """中心点索引：优先取行列几何中心 (rows//2, cols//2)；行列未知时
    退化为理论坐标最接近 (0,0) 的点。结果钳位到有效索引范围，
    防止异常调用方传入长度不足的 results 时越界崩溃。"""
    if rows and cols:
        idx = (rows // 2) * cols + (cols // 2)
        return min(idx, max(0, len(results) - 1))
    best, best_d = None, None
    for i, r in enumerate(results):
        tx, ty = r.get("theo_x"), r.get("theo_y")
        if tx is not None and ty is not None:
            d = abs(float(tx)) + abs(float(ty))
            if best_d is None or d < best_d:
                best, best_d = i, d
    if best is not None:
        return best
    return 0


def build_point_rows(results: List[dict],
                     rows: Optional[int] = None, cols: Optional[int] = None,
                     tol_upper: float = DEFAULT_TOL_UPPER,
                     tol_lower: float = DEFAULT_TOL_LOWER,
                     nominal_points: Optional[List] = None,
                     include_p1: bool = True,
                     order: str = "top_left_row_major",
                     lang: str = "zh") -> List[dict]:
    """把每点一行的检测结果重排为模板版式（每点两行）。

    results: [{"row","col","px_x","px_y","mm_x","mm_y","theo_x","theo_y","status",...}]
        mm_x/mm_y/theo_x/theo_y 约定为同一坐标系（中心归零）下的实测/名义坐标；
    nominal_points: 可选，固定名义值。
        含 P1 时 26 项 [(x,y),...]（P1 + P3..P27）：P1→[0]，P3..P27→[1..25]；
        不含 P1 时 25 项：P3..P27 直接对应 [0..24]。
        提供时 Sheet1 名义值列直接采用固定值（固定值本身就是中心基准，
        不做任何归零）；未提供时回退为 rows 中 theo_x/theo_y。
    include_p1: 是否输出 P1 中心基准行（模板版式固定含 P1，默认 True）。
    返回平铺后的点行列表：[
        {"seq":1,"name":"P1","content":"X坐标","nominal":..,"upper":..,"lower":..,"measured":..},
        {"seq":2,"name":"","content":"Y坐标",...}, ...]
    P 编号：P1=中心点基准，P2 缺失，P3 起按行主序映射行列（中心点本身
    在网格中也保有其编号，如 5×5 中心为 P15）。
    """
    c_idx = _center_index(results, rows, cols)
    has_p1 = nominal_points is not None and len(nominal_points) == len(results) + 1
    out: List[dict] = []
    seq = 0

    def emit(name: str, r: dict, fixed_nominal: Optional[tuple] = None):
        nonlocal seq
        tx, ty = r.get("theo_x"), r.get("theo_y")
        mx, my = r.get("mm_x"), r.get("mm_y")
        xy_c = _XY_CONTENT["zh" if lang != "en" else "en"]
        for idx, (content, nominal, measured) in enumerate(
                ((xy_c[0], tx, mx), (xy_c[1], ty, my))):
            seq += 1
            # 元素列仅每点首行（X坐标）填写 P 编号，Y 行留空（与模板版式一致）
            nm = name if idx == 0 else ""
            if fixed_nominal is not None:
                # 固定名义值（模板 D 列/设计网格固定值），直接采用
                nv = float(fixed_nominal[0] if idx == 0 else fixed_nominal[1])
            else:
                nv = None if nominal is None else float(nominal)
            if measured is None:
                # F 模块：失败点测量值列显式标注 "--"（不写伪造值）。
                # 兼容旧数据 status="未检出" 的语义。
                mv = "--" if _is_fail(r) else None
            else:
                mv = float(measured)
            out.append({
                "seq": seq,
                "name": nm,
                "content": content,
                "nominal": nv,
                "upper": tol_upper,
                "lower": tol_lower,
                "measured": mv,
            })

    # P1 = 中心点基准（重复输出中心点一次，与模板 P1 基准点对应）
    if include_p1:
        emit("P1", results[c_idx],
             nominal_points[0] if has_p1 else None)
    # P3..P27 按编号顺序输出（P2 缺失）；显示顺序由 order 模式决定
    disp = display_order(rows, cols, order) if (rows and cols) else list(range(len(results)))
    for i, rank in enumerate(disp):
        r = results[rank]
        if has_p1:
            fixed = nominal_points[i + 1] if i + 1 < len(nominal_points) else None
        elif nominal_points is not None:
            fixed = nominal_points[i] if i < len(nominal_points) else None
        else:
            fixed = None
        emit(f"P{i + 3}", r, fixed)
    return out


def _fmt(value) -> object:
    if value is None:
        return ""
    if isinstance(value, float):
        return round(value, 4)
    return value


def export_excel(path: str,
                 results: List[dict],
                 params: Optional[dict] = None,
                 meta: Optional[dict] = None,
                 tol_upper: Optional[float] = None,
                 tol_lower: Optional[float] = None,
                 rows: Optional[int] = None,
                 cols: Optional[int] = None,
                 nominal_points: Optional[List] = None,
                 include_p1: bool = True,
                 order: str = "top_left_row_major",
                 lang: str = "zh") -> str:
    """导出模板版式 Excel。

    results: 每点一行的检测结果 dict 列表（mm/theo 为中心归零坐标系）；
    meta:    表头元数据 12 项（key 与模板标签一致，缺省自动补空）；
    tol_upper / tol_lower: 判定公差（默认 ±0.05）；
    rows / cols: 网格行列数（用于 P 编号中心判定）；
    nominal_points: 可选，模板固定名义值 26 点（P1 + P3..P27），
        提供时 Sheet1 名义值列直接采用固定值，禁止检测反推。
    """
    tol_upper = DEFAULT_TOL_UPPER if tol_upper is None else float(tol_upper)
    tol_lower = DEFAULT_TOL_LOWER if tol_lower is None else float(tol_lower)
    meta_all = {k: "" for k in ("报表名称", "工件名称", "工件编号", "操作者", "公司名称",
                                "车间/工段", "楼层", "班别", "工序", "机台号", "客户", "时间")}
    if meta:
        meta_all.update({str(k): v for k, v in meta.items()})

    wb = Workbook()
    zh = lang != "en"

    # ---------------- Sheet1 检测数据（模板版式） ----------------
    ws = wb.active
    ws.title = _SHEET_NAMES["zh" if zh else "en"][0]
    meta_en_map = dict(zip(_META_LABELS_ZH, _META_LABELS_EN))
    def _mlabel(x: str) -> str:
        return x if zh else meta_en_map.get(x, x)
    # 表头元数据区（第 1-5 行；与模板一致：标签 A/D/F、值 B/E/G）
    for i, (a, d, f) in enumerate(_META_ROWS, 1):
        ws.cell(i, 1, f"{_mlabel(a)}:").font = Font(bold=True)
        ws.cell(i, 2, meta_all.get(a, ""))
        if d:
            ws.cell(i, 4, f"{_mlabel(d)}:").font = Font(bold=True)
            ws.cell(i, 5, meta_all.get(d, ""))
        if f:
            ws.cell(i, 6, f"{_mlabel(f)}:").font = Font(bold=True)
            ws.cell(i, 7, meta_all.get(f, ""))
    # 第 7 行表头
    headers = _HEADERS["zh" if zh else "en"]
    for j, h in enumerate(headers, 1):
        cell = ws.cell(7, j, h)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER
    # 第 8 行起：数据（每点两行）
    pts = build_point_rows(results, rows, cols, tol_upper, tol_lower,
                           nominal_points=nominal_points, include_p1=include_p1,
                           order=order, lang=lang)
    for k, p in enumerate(pts):
        rr = 8 + k
        values = [p["seq"], p["name"], p["content"],
                  _fmt(p["nominal"]), p["upper"], p["lower"], _fmt(p["measured"])]
        for j, v in enumerate(values, 1):
            cell = ws.cell(rr, j, v)
            cell.border = _BORDER
            if j in (1, 2, 3, 5, 6):
                cell.alignment = _CENTER
    for j, w in enumerate([8, 8, 12, 12, 10, 10, 14], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # ---------------- Sheet2 检测明细（保留原有逐点信息） ----------------
    ws2 = wb.create_sheet(_SHEET_NAMES["zh" if zh else "en"][1])
    headers2 = _HEADERS2["zh" if zh else "en"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
    pass_fill = PatternFill("solid", fgColor="E2EFDA")
    ng_fill = PatternFill("solid", fgColor="FCE4EC")
    fail_fill = PatternFill("solid", fgColor="EDEDED")
    for i, r in enumerate(results, 1):
        verdict = r.get("verdict", "")
        if _is_fail(r):
            # F 模块：失败行实测列统一写 "--"（与 CSV/GUI 口径一致，不写伪造值）
            px_x = px_y = mm_x = mm_y = "--"
            dx = dy = dev = ""
        else:
            px_x, px_y = _fmt(r.get("px_x")), _fmt(r.get("px_y"))
            mm_x, mm_y = _fmt(r.get("mm_x")), _fmt(r.get("mm_y"))
            dx, dy, dev = (_fmt(r.get("dx")), _fmt(r.get("dy")), _fmt(r.get("dev")))
        seq_no = display_no(i - 1, rows, cols, order)
        conf_val = r.get("conf") if not _is_fail(r) else "--"
        if conf_val is not None:
            conf_val = round(float(conf_val), 3) if not isinstance(conf_val, str) else conf_val
        ws2.append([seq_no, r.get("row", ""), r.get("col", ""),
                    px_x, px_y, mm_x, mm_y,
                    _fmt(r.get("theo_x")), _fmt(r.get("theo_y")),
                    dx, dy, dev,
                    verdict, conf_val, r.get("status", "OK")])
        if _is_fail(r):
            # F 模块：失败行浅灰底 + 斜体，与成功点视觉区分
            for c in ws2[i + 1]:
                c.fill = fail_fill
                c.font = Font(italic=True)
        elif verdict == "NG":
            for c in ws2[i + 1]:
                c.fill = ng_fill
        elif verdict == "Pass":
            for c in ws2[i + 1]:
                c.fill = pass_fill
    # ---- F 模块：统计区（失败计数独立 + 误差统计排除失败行 + 失败点行列清单）----
    fail_rows = [r for r in results if _is_fail(r)]
    ok_rows = [r for r in results if not _is_fail(r)]
    n_fail = len(fail_rows)
    n_ok = len(ok_rows)
    if n_fail > 0:
        ws2.append([])
        ws2.append(["", "", "", "", "", "", "", "", "", "", "", "",
                    ("检测失败点" if zh else "Failed points"), f"{n_fail} 个" if zh else f"{n_fail} pcs"])
        for c in ws2[ws2.max_row]:
            c.font = Font(bold=True)
    if ok_rows:
        dxs = [r["dx"] for r in ok_rows if r.get("dx") is not None]
        dys = [r["dy"] for r in ok_rows if r.get("dy") is not None]
        devs = [r["dev"] for r in ok_rows if r.get("dev") is not None]
        def _rms(vals):
            return math.sqrt(sum(v * v for v in vals) / len(vals)) if vals else 0.0
        def _mean(vals):
            return sum(vals) / len(vals) if vals else 0.0
        def _max_abs(vals):
            return max(abs(v) for v in vals) if vals else 0.0
        ws2.append(["误差统计(仅成功点,排除失败行)" if zh else "Error stats (pass only, fail excluded)", "RMS(mm)", "Max(mm)", "Mean(mm)"])
        ws2.append(["X分量" if zh else "X-Comp", round(_rms(dxs), 4), round(_max_abs(dxs), 4), round(_mean(dxs), 4)])
        ws2.append(["Y分量" if zh else "Y-Comp", round(_rms(dys), 4), round(_max_abs(dys), 4), round(_mean(dys), 4)])
        ws2.append(["合成" if zh else "Combined", round(_rms(devs), 4), round(_max_abs(devs), 4), round(_mean(devs), 4)])
    if n_fail > 0:
        ws2.append([])
        ws2.append(["失败点行列清单（便于定位图纸位置补测）" if zh else "Failed points list (for locating & re-test)"])
        ws2.append(["序号", "行", "列", "名义X(mm)", "名义Y(mm)", "状态"] if zh else ["No.", "Row", "Col", "Nominal X(mm)", "Nominal Y(mm)", "Status"])
        for i, r in enumerate(results, 1):
            if _is_fail(r):
                ws2.append([display_no(i - 1, rows, cols, order), r.get("row", ""), r.get("col", ""),
                            _fmt(r.get("theo_x")), _fmt(r.get("theo_y")),
                            r.get("status", "DETECT_FAIL")])
    for j, w in enumerate([8, 6, 6, 14, 14, 14, 14, 14, 14, 14, 14, 14, 12, 14], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w

    # ---------------- Sheet3 参数记录 ----------------
    ws3 = wb.create_sheet(_SHEET_NAMES["zh" if zh else "en"][2])
    if params:
        for k, v in params.items():
            ws3.append([str(k), str(v)])
    else:
        ws3.append(["说明", "无参数快照"] if zh else ["Note", "No parameter snapshot"])
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 40

    wb.save(path)
    return path


# ---- B8：8 列 CSV 导出（点号/名义/实测/误差/判定） ----
P0_VERDICT_THRESHOLD_MM = 0.15  # 定案 §6：0.15mm 写死，不做配置项

VERDICT_OK = "OK"
VERDICT_NG = "NG"
VERDICT_DETECT_FAIL = "DETECT_FAIL"

_CSV_HEADERS_8 = {
    "zh": ["P", "X名义", "Y名义", "X实测", "Y实测", "X误差", "Y误差", "判定"],
    "en": ["P", "X Nominal", "Y Nominal", "X Measured", "Y Measured",
           "X Error", "Y Error", "Verdict"],
}


def _verdict_8col(r: dict) -> str:
    """B8 判定列：检测失败判 DETECT_FAIL；|误差|>0.15 判 NG；否则 OK。

    误差 = 名义 - 实测；阈值 0.15mm 写死（定案 §6，不做配置项）。
    缺名义或实测数据时保守判 DETECT_FAIL，不伪造判定。
    """
    if _is_fail(r):
        return VERDICT_DETECT_FAIL
    tx, ty = r.get("theo_x"), r.get("theo_y")
    mx, my = r.get("mm_x"), r.get("mm_y")
    if tx is None or ty is None or mx is None or my is None:
        return VERDICT_DETECT_FAIL
    ex = float(tx) - float(mx)
    ey = float(ty) - float(my)
    if abs(ex) <= P0_VERDICT_THRESHOLD_MM and abs(ey) <= P0_VERDICT_THRESHOLD_MM:
        return VERDICT_OK
    return VERDICT_NG


def build_csv_rows(results: List[dict],
                   rows: Optional[int] = None, cols: Optional[int] = None,
                   nominal_points: Optional[List] = None,
                   order: str = "bottom_left_row_major",
                   lang: str = "zh") -> List[dict]:
    """B8：每点一行 8 列（P / X名义 / Y名义 / X实测 / Y实测 / X误差 / Y误差 / 判定）。

    - 名义值固定（设计值）：优先 nominal_points（与 results 同序逐点覆盖），
      否则取 results 的 theo_x/theo_y；
    - 实测值 = mm_x/mm_y（调用方保证与名义同坐标系：中心归零）；
    - 误差 = 名义 - 实测（定案 §5）；判定阈值 0.15mm 写死；
    - 排序复用 point_order.display_order，4 种起始角，默认左下。
    """
    if rows and cols:
        order_ranks = display_order(rows, cols, order)
    else:
        order_ranks = list(range(len(results)))
    out: List[dict] = []
    for rank in order_ranks:
        r = results[rank]
        p_no = display_no(rank, rows, cols, order) if (rows and cols) else rank + 1
        if nominal_points is not None and rank < len(nominal_points):
            nx, ny = nominal_points[rank][0], nominal_points[rank][1]
        else:
            nx, ny = r.get("theo_x"), r.get("theo_y")
        mx, my = r.get("mm_x"), r.get("mm_y")
        ex = None if (_is_fail(r) or nx is None or mx is None) else float(nx) - float(mx)
        ey = None if (_is_fail(r) or ny is None or my is None) else float(ny) - float(my)
        out.append({
            "P": p_no,
            "X名义": nx, "Y名义": ny,
            "X实测": mx, "Y实测": my,
            "X误差": ex, "Y误差": ey,
            "判定": _verdict_8col(r),
        })
    return out


def export_csv(path: str,
               results: List[dict],
               rows: Optional[int] = None, cols: Optional[int] = None,
               nominal_points: Optional[List] = None,
               order: str = "bottom_left_row_major",
               lang: str = "zh") -> str:
    """B8：导出 8 列 CSV（utf-8-sig，Excel 可直接打开），返回 path。"""
    import csv
    headers = _CSV_HEADERS_8["zh" if lang != "en" else "en"]
    pts = build_csv_rows(results, rows, cols, nominal_points, order, lang)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for p in pts:
            w.writerow([p["P"], _fmt(p["X名义"]), _fmt(p["Y名义"]),
                        _fmt(p["X实测"]), _fmt(p["Y实测"]),
                        _fmt(p["X误差"]), _fmt(p["Y误差"]), p["判定"]])
    return path


