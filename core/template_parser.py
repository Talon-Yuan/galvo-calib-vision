# -*- coding: utf-8 -*-
"""质检模板解析器：从模板工作簿（如 工作簿1.xlsx）读取导出版式所需的
判定公差、表头元数据与固定名义值，供导出层与主窗口判定口径使用。

模板 Sheet1 约定（与 工作簿1.xlsx 一致）：
- 第 1-5 行：表头元数据区，表单式布局，标签在 A/D/F 列（带冒号），值在 B/E/G 列；
- 第 6 行：空行；
- 第 7 行：检测数据表表头（序号/元素/内容/名义值/上公差/下公差/测量值）；
- 第 8 行起：每点拆两行（内容=X坐标/Y坐标），名义值在第 4 列(D)，上/下公差在第 5/6 列。

名义值是用户以二次元测量仪实测得到的固定设计值，软件不得用检测拟合坐标
反推或替换名义值——本解析器将其原样提取为 nominal_points 供上层使用。
"""
from typing import Dict, List, Tuple

from openpyxl import load_workbook

# 表头元数据 12 项（与模板标签一致，标签可能带冒号，解析时去除）
META_LABELS = [
    "报表名称", "工件名称", "工件编号", "操作者", "公司名称",
    "车间/工段", "楼层", "班别", "工序", "机台号", "客户", "时间",
]


def _clean_label(text) -> str:
    """去掉标签文本末尾的冒号（中英文）。"""
    if text is None:
        return ""
    return str(text).strip().rstrip(":：")


def parse_template(path: str) -> Dict:
    """解析模板工作簿，返回：
    {
      "meta": {12 项元数据标签: 值},
      "tol_upper": float,              # 上公差（正数，如 0.05）
      "tol_lower": float,              # 下公差（负数，如 -0.05）
      "nominal_points": List[Tuple[float, float]],  # 固定名义值 (X,Y)，26 点 = P1 + P3..P27
      "nominal_names": List[str],      # 与 nominal_points 对应的点名称
      "path": 模板路径,
    }
    """
    wb = load_workbook(path, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active

    meta: Dict[str, object] = {}
    # 第 1-5 行扫描 A/D/F 标签 + B/E/G 值（模板真实布局：C 列全空）
    for row in range(1, 6):
        for label_col, val_col in ((1, 2), (4, 5), (6, 7)):
            label = _clean_label(ws.cell(row, label_col).value)
            if not label:
                continue
            value = ws.cell(row, val_col).value
            meta[label] = "" if value is None else value

    # 检测数据表头行：扫描前 15 行找含"序号"的表头行
    header_row = 7
    for row in range(1, 16):
        if str(ws.cell(row, 1).value or "").strip() == "序号":
            header_row = row
            break
    data_first = header_row + 1
    tol_upper = DEFAULT_TOL_UPPER = 0.05
    tol_lower = DEFAULT_TOL_LOWER = -0.05
    upper_cell = ws.cell(data_first, 5).value
    lower_cell = ws.cell(data_first, 6).value
    if isinstance(upper_cell, (int, float)):
        tol_upper = float(upper_cell)
    if isinstance(lower_cell, (int, float)):
        tol_lower = float(lower_cell)

    # 解析 D 列固定名义值：内容列(C)为 X坐标/Y坐标 交替，每点两行合成 (X, Y)。
    # 名义值直接取自模板固定值（用户二次元测量结果），严禁用检测坐标反推替换。
    nominal_points: List[Tuple[float, float]] = []
    nominal_names: List[str] = []
    row = data_first
    while row <= ws.max_row:
        content = str(ws.cell(row, 3).value or "").strip()
        nominal = ws.cell(row, 4).value
        if not isinstance(nominal, (int, float)) or content not in ("X坐标", "Y坐标"):
            break
        if content == "X坐标":
            nominal_names.append(str(ws.cell(row, 2).value or "").strip())
            nominal_points.append([float(nominal), 0.0])
        else:  # Y坐标，闭合当前点
            if nominal_points:
                nominal_points[-1][1] = float(nominal)
        row += 1
    nominal_points = [tuple(p) for p in nominal_points]

    return {
        "meta": meta,
        "tol_upper": tol_upper,
        "tol_lower": tol_lower,
        "nominal_points": nominal_points,
        "nominal_names": nominal_names,
        "path": path,
    }
